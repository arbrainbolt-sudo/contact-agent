"""Events page — a blueprint you can drop into the existing app.

Adds a /events page with five boxes:

    5. My Interest        (anything you've ticked, pulled to the top)
    1. Philadelphia
    2. NY / NJ / DE
    3. All over USA
    4. Large tech companies

Data lives in an "Events" worksheet inside the same results.xlsx that
store.py manages. The big-tech list is seeded automatically the first
time the sheet is created.

To wire it in, add two lines to app.py:

    from events import bp as events_bp          # near the other imports
    app.register_blueprint(events_bp)           # after app = Flask(__name__)

Then add a link to your nav bar:

    <a href="/events">Events</a>
"""

import os
import uuid
import threading

from flask import Blueprint, request, redirect, url_for, render_template_string
from openpyxl import Workbook, load_workbook

import store

bp = Blueprint("events", __name__)

SHEET = "Events"
HEADERS = [
    "event_id",
    "category",
    "date",
    "company",
    "event",
    "location",
    "meant_for",
    "register_url",
    "interested",
    "notes",
]

CATEGORIES = [
    ("philadelphia", "1. Philadelphia"),
    ("tristate", "2. New York, New Jersey, Delaware"),
    ("usa", "3. All over USA"),
    ("bigtech", "4. Large tech companies"),
]
CATEGORY_LABELS = dict(CATEGORIES)

# Suggestions offered in the "Meant for" box. Free text — these are just hints.
MEANT_FOR_SUGGESTIONS = [
    "Hackathon / build night",
    "Tech talk + Q&A",
    "Hands-on workshop",
    "Networking meetup",
    "Career meetup",
    "User group",
    "Open-source community",
    "Tech roundtable",
    "Founders / investors",
    "Fintech",
    "Startup / entrepreneurship",
]

_lock = threading.Lock()


# --------------------------------------------------------------- seed data
#
# Dates below are the month these events USUALLY land in, not confirmed 2027
# dates — most 2027 dates aren't published yet. Edit them on the page as
# each one is announced.

BIGTECH_SEED = [
    # company, event, typical timing, usual location, focus, url
    ("Apple", "WWDC", "Typically June", "Cupertino, CA + online",
     "iOS, macOS, Swift, Apple Intelligence", "https://developer.apple.com/wwdc/"),
    ("Google", "Google I/O", "Typically May", "Mountain View, CA + online",
     "Android, Gemini/AI, Chrome, dev platforms", "https://io.google/"),
    ("Google Cloud", "Google Cloud Next", "Typically April", "Las Vegas, NV",
     "Cloud, Gemini, data, enterprise AI", "https://cloud.withgoogle.com/next"),
    ("Microsoft", "Microsoft Build", "Typically May", "Seattle, WA + online",
     "Azure, Windows, Copilot, AI agents", "https://build.microsoft.com/"),
    ("Microsoft", "Microsoft Ignite", "Typically November", "USA (varies)",
     "Enterprise IT, Azure, security, M365", "https://ignite.microsoft.com/"),
    ("AWS", "AWS re:Invent", "Typically early December", "Las Vegas, NV",
     "Cloud, AI, databases, serverless", "https://reinvent.awsevents.com/"),
    ("AWS", "AWS Summit New York", "Typically July", "New York, NY",
     "Regional AWS cloud/developer event", "https://aws.amazon.com/events/summits/"),
    ("NVIDIA", "NVIDIA GTC", "Typically March", "San Jose, CA",
     "AI, GPUs, robotics, accelerated computing", "https://www.nvidia.com/gtc/"),
    ("OpenAI", "OpenAI DevDay", "Varies (autumn)", "San Francisco, CA",
     "Models, APIs, agents, developer tools", "https://openai.com/devday/"),
    ("Meta", "Meta Connect", "Typically September", "Menlo Park, CA",
     "AI, smart glasses, VR/AR, Horizon", "https://www.meta.com/connect/"),
    ("GitHub", "GitHub Universe", "Typically October", "San Francisco, CA",
     "Copilot, AI coding, DevOps", "https://githubuniverse.com/"),
    ("Salesforce", "Dreamforce", "Typically October", "San Francisco, CA",
     "Agentforce/AI, CRM, enterprise apps", "https://www.salesforce.com/dreamforce/"),
    ("Oracle", "Oracle AI World", "Typically October", "Las Vegas, NV",
     "Oracle Cloud, databases, enterprise AI", "https://www.oracle.com/events/"),
    ("SAP", "SAP Sapphire", "Typically May", "Orlando, FL",
     "ERP, business AI, enterprise transformation", "https://www.sap.com/events/sapphire.html"),
    ("IBM", "IBM Think", "Typically May", "USA (varies)",
     "Enterprise AI, hybrid cloud, strategy", "https://www.ibm.com/events/"),
    ("IBM", "IBM TechXchange", "Typically October", "Las Vegas, NV",
     "Technical / developer-focused IBM event", "https://www.ibm.com/events/"),
    ("Cisco", "Cisco Live US", "Typically June", "USA (varies)",
     "Networking, security, AI, observability", "https://www.ciscolive.com/"),
    ("Cisco", "WebexOne", "Typically autumn", "USA (varies)",
     "Collaboration, contact centre, AI comms", "https://www.webexone.com/"),
    ("ServiceNow", "Knowledge", "Typically May", "Las Vegas, NV",
     "Workflows, enterprise automation, AI agents", "https://www.servicenow.com/events/knowledge.html"),
    ("Snowflake", "Snowflake Summit", "Typically June", "San Francisco, CA",
     "Data cloud, analytics, AI", "https://www.snowflake.com/summit/"),
    ("Databricks", "Data + AI Summit", "Typically June", "San Francisco, CA",
     "Lakehouse, data engineering, GenAI, ML", "https://www.databricks.com/dataaisummit"),
    ("Red Hat", "Red Hat Summit", "Typically May", "USA (varies)",
     "Linux, OpenShift, Kubernetes, hybrid cloud", "https://www.redhat.com/en/summit"),
    ("Splunk", ".conf", "Typically September", "USA (varies)",
     "Observability, security, SIEM, IT ops", "https://conf.splunk.com/"),
    ("Atlassian", "Team", "Typically April", "Anaheim / Las Vegas",
     "Jira, Confluence, dev productivity, AI", "https://www.atlassian.com/team"),
    ("Adobe", "Adobe MAX", "Typically October", "Los Angeles, CA",
     "Creative tech, generative AI, Firefly", "https://max.adobe.com/"),
    ("Adobe", "Adobe Summit", "Typically March", "Las Vegas, NV",
     "Digital experience, martech, AI", "https://summit.adobe.com/"),
    ("Dell Technologies", "Dell Technologies World", "Typically May", "Las Vegas, NV",
     "Infrastructure, storage, AI, enterprise IT", "https://www.dell.com/en-us/dt/events/"),
    ("Broadcom / VMware", "VMware Explore", "Typically August", "Las Vegas, NV",
     "Virtualisation, private cloud, infrastructure", "https://www.vmware.com/explore"),
    ("Intel", "Intel developer events", "Varies", "USA (varies)",
     "CPUs, AI, edge, developer technology", "https://www.intel.com/content/www/us/en/events/"),
    ("AMD", "Advancing AI", "Varies", "USA (varies)",
     "CPUs, GPUs, AI infrastructure", "https://www.amd.com/en/corporate/events.html"),
    ("Qualcomm", "Snapdragon Summit", "Typically autumn", "Maui, HI",
     "Mobile chips, PCs, edge AI", "https://www.qualcomm.com/snapdragon/summit"),
    ("Samsung", "Samsung Developer Conference", "Typically October", "San Jose, CA",
     "Mobile, SmartThings, devices, AI", "https://developer.samsung.com/sdc"),
]


# ------------------------------------------------------------------ storage

def _backup():
    """Use store's backup if it exposes one; never let a backup failure block."""
    try:
        fn = getattr(store, "backup_now", None) or getattr(store, "_backup", None)
        if fn:
            fn()
    except Exception as exc:
        print(f"[events] backup skipped: {exc}")


def _workbook():
    if os.path.exists(store.XLSX_FILE):
        return load_workbook(store.XLSX_FILE)
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def _sheet(wb):
    """Return the Events worksheet, creating and seeding it if it's missing."""
    if SHEET in wb.sheetnames:
        return wb[SHEET], False
    ws = wb.create_sheet(SHEET)
    ws.append(HEADERS)
    for company, event, when, where, focus, url in BIGTECH_SEED:
        ws.append([
            uuid.uuid4().hex[:8], "bigtech", when, company,
            event, where, focus, url, "", "",
        ])
    return ws, True


def _text(v):
    return "" if v is None else str(v).replace("_x000D_", "").strip()


def all_events():
    with _lock:
        wb = _workbook()
        ws, created = _sheet(wb)
        if created:
            _backup()
            wb.save(store.XLSX_FILE)
        rows = []
        for raw in ws.iter_rows(min_row=2, values_only=True):
            if not raw or not any(raw):
                continue
            row = dict(zip(HEADERS, [_text(c) for c in raw]))
            row["interested"] = row["interested"].lower() in ("1", "yes", "true", "x")
            rows.append(row)
        return rows


def _rewrite(rows):
    _backup()
    wb = _workbook()
    if SHEET in wb.sheetnames:
        wb.remove(wb[SHEET])
    ws = wb.create_sheet(SHEET)
    ws.append(HEADERS)
    for row in rows:
        ws.append([
            row["event_id"], row["category"], row["date"], row["company"],
            row["event"], row["location"], row["meant_for"],
            row["register_url"], "yes" if row["interested"] else "", row["notes"],
        ])
    if not wb.sheetnames:
        wb.create_sheet("Sheet1")
    wb.save(store.XLSX_FILE)


def add_event(data):
    rows = all_events()
    rows.append({
        "event_id": uuid.uuid4().hex[:8],
        "category": data.get("category", "philadelphia"),
        "date": data.get("date", ""),
        "company": data.get("company", ""),
        "event": data.get("event", ""),
        "location": data.get("location", ""),
        "meant_for": data.get("meant_for", ""),
        "register_url": data.get("register_url", ""),
        "interested": False,
        "notes": data.get("notes", ""),
    })
    with _lock:
        _rewrite(rows)


def set_interest(event_id, wanted):
    rows = all_events()
    for row in rows:
        if row["event_id"] == event_id:
            row["interested"] = wanted
    with _lock:
        _rewrite(rows)


def delete_event(event_id):
    rows = [r for r in all_events() if r["event_id"] != event_id]
    with _lock:
        _rewrite(rows)


# ------------------------------------------------------------------- routes

@bp.route("/events")
def events_page():
    rows = all_events()
    interested = [r for r in rows if r["interested"]]
    boxes = [
        (key, label, [r for r in rows if r["category"] == key])
        for key, label in CATEGORIES
    ]
    return render_template_string(
        PAGE,
        interested=interested,
        boxes=boxes,
        categories=CATEGORIES,
        suggestions=MEANT_FOR_SUGGESTIONS,
        labels=CATEGORY_LABELS,
        workbook=store.XLSX_FILE,
        total=len(rows),
    )


@bp.route("/events/toggle/<event_id>", methods=["POST"])
def events_toggle(event_id):
    set_interest(event_id, request.form.get("interested") == "on")
    return redirect(url_for("events.events_page"))


@bp.route("/events/add", methods=["POST"])
def events_add():
    if request.form.get("event", "").strip() or request.form.get("company", "").strip():
        add_event(request.form)
    return redirect(url_for("events.events_page"))


@bp.route("/events/delete/<event_id>", methods=["POST"])
def events_delete(event_id):
    delete_event(event_id)
    return redirect(url_for("events.events_page"))


# ------------------------------------------------------------------ template

PAGE = """
<!doctype html>
<html><head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Events</title>
<style>
  body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
         margin: 0; padding: 20px; background: #f6f7f9; color: #1c1c1c; }
  .wrap { max-width: 1400px; margin: 0 auto; }
  h1 { font-size: 24px; margin: 0 0 4px; color: #1a3a6b; }
  .sub { font-size: 13px; color: #666; margin-bottom: 18px; }
  nav { margin-bottom: 18px; display: flex; gap: 16px; flex-wrap: wrap; font-size: 14px; }
  nav a { color: #1a3a6b; text-decoration: none; }
  nav a:hover { text-decoration: underline; }

  .box { background: #fff; border: 1px solid #e0e3e8; border-radius: 10px;
         padding: 16px; margin-bottom: 18px; }
  .box h2 { font-size: 15px; margin: 0 0 12px; text-transform: uppercase;
            letter-spacing: .04em; color: #444; }
  .box.mine { border: 2px solid #96C950; background: #fbfdf6; }
  .box.mine h2 { color: #4d7a1f; }
  .grid { display: grid; grid-template-columns: repeat(auto-fit, minmax(440px, 1fr)); gap: 18px; }

  table { width: 100%; border-collapse: collapse; font-size: 13px; }
  th { text-align: left; font-size: 11px; text-transform: uppercase; letter-spacing: .04em;
       color: #888; border-bottom: 1px solid #e6e8ec; padding: 6px 8px; font-weight: 600; }
  td { padding: 8px; border-bottom: 1px solid #f0f1f4; vertical-align: top; }
  tr:last-child td { border-bottom: none; }
  td.co { font-weight: 600; }
  td.dt { white-space: nowrap; color: #555; }
  .tag { display: inline-block; font-size: 11px; background: #eef1f6; color: #4a5568;
         border-radius: 4px; padding: 1px 6px; margin-right: 4px; }
  .empty { color: #999; font-size: 13px; padding: 10px 2px; }
  a.reg { color: #1a3a6b; }
  input[type=checkbox] { width: 17px; height: 17px; }
  .del { border: none; background: none; color: #bbb; cursor: pointer; font-size: 15px; }
  .del:hover { color: #c0392b; }

  details { margin-top: 12px; }
  summary { cursor: pointer; font-size: 13px; color: #1a3a6b; }
  .form { display: grid; grid-template-columns: repeat(auto-fit, minmax(150px, 1fr));
          gap: 8px; margin-top: 10px; }
  .form input, .form select { padding: 7px 9px; border: 1px solid #d5d9e0;
                              border-radius: 6px; font-size: 13px; width: 100%; }
  .btn { background: #1a3a6b; color: #fff; border: none; border-radius: 6px;
         padding: 8px 16px; font-size: 13px; cursor: pointer; }
  @media (max-width: 760px) {
    .grid { grid-template-columns: 1fr; }
    body { padding: 12px; }
  }
</style>
</head><body><div class="wrap">

<nav>
  <a href="/">Contacts</a>
  <a href="/news">News</a>
  <a href="/crm">CRM</a>
  <a href="/events"><b>Events</b></a>
</nav>

<h1>Events</h1>
<div class="sub">{{ total }} events &middot; saved in {{ workbook }} on the "Events" sheet</div>

{% macro event_table(rows, show_cat=False) %}
  {% if rows %}
  <table>
    <tr>
      <th style="width:28px"></th>
      <th>Date</th><th>Company</th><th>Event</th><th>Location</th>
      <th>Meant for</th><th>Register</th>{% if show_cat %}<th>Where</th>{% endif %}<th></th>
    </tr>
    {% for r in rows %}
    <tr>
      <td>
        <form method="post" action="/events/toggle/{{ r.event_id }}" style="margin:0">
          <input type="checkbox" name="interested" {% if r.interested %}checked{% endif %}
                 onchange="this.form.submit()">
        </form>
      </td>
      <td class="dt">{{ r.date or "—" }}</td>
      <td class="co">{{ r.company }}</td>
      <td>{{ r.event }}</td>
      <td>{{ r.location }}</td>
      <td>{% if r.meant_for %}<span class="tag">{{ r.meant_for }}</span>{% endif %}</td>
      <td>{% if r.register_url %}<a class="reg" href="{{ r.register_url }}" target="_blank"
            rel="noopener">Register</a>{% endif %}</td>
      {% if show_cat %}<td><span class="tag">{{ labels[r.category] }}</span></td>{% endif %}
      <td>
        <form method="post" action="/events/delete/{{ r.event_id }}" style="margin:0"
              onsubmit="return confirm('Delete this event?')">
          <button class="del" type="submit" title="Delete">&times;</button>
        </form>
      </td>
    </tr>
    {% endfor %}
  </table>
  {% else %}
  <div class="empty">Nothing here yet.</div>
  {% endif %}
{% endmacro %}

{% macro add_form(cat) %}
<details>
  <summary>+ Add an event</summary>
  <form class="form" method="post" action="/events/add">
    <input type="hidden" name="category" value="{{ cat }}">
    <input name="date" placeholder="Date e.g. 14 Oct 2026">
    <input name="company" placeholder="Company / organiser">
    <input name="event" placeholder="Event name">
    <input name="location" placeholder="Location">
    <input name="meant_for" list="meantfor" placeholder="Meant for">
    <input name="register_url" placeholder="Registration link">
    <button class="btn" type="submit">Add</button>
  </form>
</details>
{% endmacro %}

<!-- Box 5, pulled to the top -->
<div class="box mine">
  <h2>&#9733; My Interest &mdash; {{ interested|length }}</h2>
  {{ event_table(interested, show_cat=True) }}
</div>

<div class="grid">
  {% for key, label, rows in boxes %}
  <div class="box">
    <h2>{{ label }} &mdash; {{ rows|length }}</h2>
    {{ event_table(rows) }}
    {{ add_form(key) }}
  </div>
  {% endfor %}
</div>

<datalist id="meantfor">
  {% for s in suggestions %}<option value="{{ s }}">{% endfor %}
</datalist>

</div></body></html>
"""